"""Order book in the browser: customer + price list, order form (codes, teeth -> quantity, surcharge), search,
redo (R), KTV / QC workflow and team chat (KTV never sees money), payment, reports with chart, Excel export
(checked with openpyxl), Excel import, 3D link with preview picture, phone layout. Screenshots in the folder given.
Run against: PAY=1 node dev/server.mjs public 8788 mat-khau-thu"""
import sys, os, time, json, re
from playwright.sync_api import sync_playwright
import openpyxl
B = os.environ.get("D8_BASE", "http://localhost:8788")
OUT = sys.argv[1] if len(sys.argv) > 1 else "shots-orders"
THEME = sys.argv[2] if len(sys.argv) > 2 else "light"
os.makedirs(OUT, exist_ok=True)
FAILS, errors = [], []
def ok(c, m):
    print(("PASS " if c else "FAIL ") + m)
    if not c: FAILS.append(m)
def block(route):
    u = route.request.url
    if "fonts.g" in u or "accounts.google" in u or "cdnjs" in u: return route.abort()
    return route.continue_()
SIGNIN = """async ([email,name]) => { const tok = await (await fetch('/__test/idtoken?email='+encodeURIComponent(email)+'&name='+encodeURIComponent(name))).text();
  const r = await fetch('/api/auth/google', {method:'POST', headers:{'Content-Type':'application/json','X-D8':'1'}, body: JSON.stringify({credential: tok})}); return r.status; }"""
API = """async ([path, body, method]) => { const r = await fetch(path, {method: method || (body ? 'POST' : 'GET'), headers:{'Content-Type':'application/json','X-D8':'1'}, body: body ? JSON.stringify(body) : undefined}); return [r.status, await r.json().catch(()=>({}))]; }"""
# nothing in the chat thread is wider than the thread, and the thread does not scroll sideways
THREAD_FITS = """e => { const r = e.getBoundingClientRect(); return e.scrollWidth <= e.clientWidth + 1 &&
  [...e.querySelectorAll('*')].every((c) => { const q = c.getBoundingClientRect(); return !q.width || q.right <= r.right + 1; }); }"""
stamp = str(int(time.time()))[-6:]
CODE = "UI" + stamp
KTV, QC = f"ktv{stamp}@d8.example", f"qc{stamp}@d8.example"
EXO = os.path.join(OUT, "scene-exocad.html")
open(EXO, "w").write("""<!doctype html><html><head><title>exocad scene</title></head><body style="margin:0;background:#473A6D">
<canvas id="c" width="1200" height="630" style="display:block;width:100vw;height:100vh"></canvas>
<script>window.DentalWebGL={hasEmbeddedData:function(){return true},m_SceneManager:{m_MeshSort:{m_SortedMeshs:[{IsLoaded:true}]}}};
var g=document.getElementById('c').getContext('webgl'); if (g) { g.clearColor(0.86,0.74,0.52,1); g.clear(g.COLOR_BUFFER_BIT); }</script></body></html>""")

with sync_playwright() as p:
    b = p.chromium.launch()
    def ctx(vp=None, mobile=False):
        c = b.new_context(viewport=vp or {"width": 1366, "height": 900}, is_mobile=mobile, has_touch=mobile, reduced_motion="reduce", accept_downloads=True)
        c.route("**/*", block)
        c.add_init_script(f"try{{localStorage.setItem('d8-theme','{THEME}')}}catch(e){{}}")
        return c
    def watch(pg, name):
        pg.on("pageerror", lambda e: errors.append(name + " PAGEERROR " + str(e)))
        pg.on("console", lambda m: errors.append(name + " console: " + m.text) if m.type == "error" and "ERR_FAILED" not in m.text and "401" not in m.text and "403" not in m.text and "404" not in m.text and "409" not in m.text else None)
    def text(pg, sel): return pg.inner_text(sel) if pg.query_selector(sel) else ""

    oc = ctx(); op = oc.new_page(); watch(op, "owner")
    op.goto(B + "/admin"); op.evaluate(SIGNIN, ["d.eightdesigncenter@gmail.com", "Danny Do"])
    op.evaluate(API, ["/api/products/sample", {}])
    for e, role in [(KTV, "designer"), (QC, "qc")]: op.evaluate(API, ["/api/staff", {"email": e, "role": role}])
    op.reload(); op.wait_for_timeout(1200)

    # customer through the form, then its price list
    op.goto(B + "/admin#customers/new"); op.wait_for_timeout(700)
    ok(op.query_selector("input#cfCode") is None, "customer code is not typed in")
    op.select_option("#cfKind", "lab"); op.wait_for_timeout(700)
    lab_code = op.inner_text("#cfCode")
    op.select_option("#cfKind", "doctor"); op.wait_for_timeout(700)
    dr_code = op.inner_text("#cfCode")
    ok(re.match(r"^LAB\d{6,}$", lab_code) and re.match(r"^DR\d{6,}$", dr_code) and lab_code[3:] == dr_code[2:], "code shown from the kind: " + lab_code + " / " + dr_code)
    op.fill("#cfName", "Lux Smile Studio " + stamp); op.select_option("#cfCur", "USD")
    op.fill("#cfTerms", "Trả từng ca qua Payoneer"); op.fill("#cfNote", "Thích tiếp xúc chặt, gửi file trước 9h"); op.click("#cfForm [type=submit]"); op.wait_for_timeout(1500)
    m = re.search(r"#cust/(\d+)", op.url); ok(m is not None, "customer created, customer page opened")
    CID = m.group(1) if m else "0"
    st, cd = op.evaluate(API, [f"/api/customers/{CID}"])
    CODE = cd.get("customer", {}).get("code", CODE)
    ok(CODE == dr_code and cd["customer"].get("note", "").startswith("Thích tiếp xúc"), "customer saved as " + CODE + " with the internal note")
    op.fill('#cuPrices [data-p="CR"]', "35"); op.click("#cuPrices [data-save]"); op.wait_for_timeout(900)
    st, pr = op.evaluate(API, [f"/api/prices?customer={CID}"])
    ok(any(x["product"] == "CR" and x["price"] == 35 for x in pr.get("own", [])), "customer price saved from the customer page")
    op.screenshot(path=f"{OUT}/customer.png", full_page=True)

    # price lists in the owner's layout (made up prices): list price, level 2; a customer on level 2
    def price_file(path, prices, quote="Quote on request"):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "PRICE LIST"
        ws.append(["DENTAL CAD DESIGN SERVICES"]); ws.append(["STANDARD PRICE LIST"]); ws.append(["Effective from: test | Currency: USD"]); ws.append([]); ws.append([])
        ws.append(["Prepared for: ....", None, None, "Quote ref:", "Q-TEST"]); ws.append([])
        ws.append(["No.", "SERVICE", "DESCRIPTION", "UNIT", "UNIT PRICE (USD)"])
        ws.append(["CROWN & BRIDGE"]); ws.append([7, "Full Contour " + stamp, "Mão toàn phần " + stamp, "Tooth", prices[0]]); ws.append([10, "Full Cantilever Bridge " + stamp, "Cầu Cantilever " + stamp, "Bridge", prices[1]])
        ws.append(["ORTHODONTICS / ALIGNER"]); ws.append([48, "Visualizing Ortho Analyzer " + stamp, "Trực quan hoá " + stamp, "Case", quote])
        ws.append([]); ws.append(["All prices are per unit stated."])
        wb.save(path)
    fmax, fmid = os.path.join(OUT, "BANG_GIA_MAX_test.xlsx"), os.path.join(OUT, "BANG_GIA_MID_test.xlsx")
    price_file(fmax, [5.5, 9.5]); price_file(fmid, [4.5, 8.0])
    op.goto(B + "/admin#catalog"); op.wait_for_timeout(1500)
    ok("Giá mặc định" not in text(op, "#p-catalog"), "no default USD / VND price columns")
    op.click('[data-plimp=""]'); op.wait_for_timeout(300)
    op.set_input_files("#plBox [data-file]", fmax); op.wait_for_timeout(2200)
    summ = text(op, "#plBox .pl-sum").lower()
    ok("3 dịch vụ" in summ and "2 có giá" in summ and "1 báo giá theo ca" in summ, "price file preview: 3 services, 2 priced, 1 quote on request (" + summ.replace("\n", " ") + ")")
    codes = op.eval_on_selector_all("#plBox .pl-prev tbody td:nth-child(2) b", "els => els.map((e) => e.textContent)")
    op.screenshot(path=f"{OUT}/pricelist-preview.png", full_page=True)
    op.click("#plBox [data-go]"); op.wait_for_timeout(2200)
    st, pl = op.evaluate(API, ["/api/products"])
    new = {p["name_en"]: p for p in pl.get("products", []) if p["name_en"].endswith(stamp)}
    ok(len(new) == 3 and all(new[k]["code"] in codes for k in new) and new["Full Cantilever Bridge " + stamp]["unit"] == "bridge", "codes created from the file: " + ", ".join(sorted(p["code"] for p in new.values())))
    op.click('[data-plimp="mid"]'); op.wait_for_timeout(300)
    op.set_input_files("#plBox [data-file]", fmid); op.wait_for_timeout(2200)
    ok("mã mới" not in text(op, "#plBox .pl-sum").lower(), "level 2 file matches the same codes")
    op.click("#plBox [data-go]"); op.wait_for_timeout(2200)
    ok(re.search(r"\d+ mã có giá", text(op, ".pl-tile:nth-child(2)")) is not None, "price level tiles count priced codes")
    op.screenshot(path=f"{OUT}/catalog.png", full_page=True)
    FC = new["Full Contour " + stamp]["code"]
    op.goto(B + f"/admin#cust/{CID}"); op.wait_for_timeout(1500)
    op.select_option("#cuTier", "mid"); op.wait_for_timeout(1600)
    ok(op.get_attribute(f'#cuPrices [data-p="{FC}"]', "placeholder") == "4.5", "level 2 price shown for the customer")
    st, o = op.evaluate(API, ["/api/orders", {"customer_id": int(CID), "received": "2026-10-22", "items": [{"product": FC, "qty": 2}]}])
    ok(o.get("order", {}).get("total") == 9, "order for a level 2 customer: 2 × 4.5")
    if o.get("order"): op.evaluate("async (id) => (await fetch('/api/orders/' + id, {method: 'DELETE', headers: {'X-D8': '1'}})).status", o["order"]["id"])
    op.select_option("#cuTier", ""); op.wait_for_timeout(1500)

    # new order through the form
    op.goto(B + f"/admin#orders/new?customer={CID}"); op.wait_for_timeout(1500)
    ok(CODE in op.input_value("#ofCust"), "form opens with the customer filled in")
    code_preview = op.inner_text("#ofCodeOut")
    ok(re.match(r"^\d{8,9}$", code_preview) is not None, "code preview " + code_preview)
    op.fill("#ofClinic", "Lux Smile"); op.fill("#ofDoctor", "Dr Kiệt"); op.fill("#ofPatient", "Chuang Tsung Wei"); op.fill("#ofRef", "KN" + stamp)
    op.fill('#ofItems [data-i="0"] [data-k="product"]', "CR"); op.wait_for_timeout(500)
    op.fill('#ofItems [data-i="0"] [data-k="teeth"]', "11-13"); op.wait_for_timeout(700)
    ok(op.input_value('#ofItems [data-i="0"] [data-k="qty"]') == "3", "teeth 11-13 → quantity 3")
    op.click("#ofAdd"); op.fill('#ofItems [data-i="1"] [data-k="product"]', "IMC"); op.wait_for_timeout(600)
    op.fill('#ofItems [data-i="1"] [data-k="price"]', "45"); op.wait_for_timeout(600)
    op.click('[data-pct="50"]'); op.wait_for_timeout(300)
    ok("US$225" in text(op, "#ofMoney"), "form total (3×35 + 45) +50% = US$225: " + text(op, "#ofMoney").replace("\n", " "))
    op.select_option("#ofDes", KTV); op.select_option("#ofQc", QC)
    op.fill("#ofChat", "Lưu ý tiếp xúc nhẹ, khớp cắn 0.4 mm")
    op.screenshot(path=f"{OUT}/order-form.png", full_page=True)
    # the summary panel never spills over, even where the text renders larger (Windows fonts, minimum font size)
    SPILL = """() => { const c = document.querySelector('.of-side .card'), cr = c.getBoundingClientRect(), pad = parseFloat(getComputedStyle(c).paddingRight);
      return [...c.querySelectorAll('*')].filter((e) => { const r = e.getBoundingClientRect(); return r.width && r.right > cr.right - pad + 0.5; }).map((e) => e.id || e.tagName).slice(0, 5); }"""
    ok(op.evaluate(SPILL) == [], "order form panel: nothing spills over the card")
    big = op.add_style_tag(content=".seg span{font-size:12.5px !important} .of-side .input{font-size:17px !important}"); op.wait_for_timeout(200)
    ok(op.evaluate(SPILL) == [], "order form panel: nothing spills over with larger text either")
    big.evaluate("e => e.remove()")
    op.click("#ofSave"); op.wait_for_timeout(2000)
    m = re.search(r"#orders/(\d+)$", op.url); ok(m is not None, "order saved, order page opened")
    OID = m.group(1) if m else "0"
    band = op.inner_text(".pt-band")
    ok(code_preview in band and "Chuang Tsung Wei" in band, "order page title: code + patient")
    body = op.inner_text("#p-order")
    ok("US$225" in body and "Đang thiết kế" in body, "order page: total and status")
    op.wait_for_timeout(600)
    ok("khớp cắn 0.4" in text(op, "#oChat"), "first team message is in the order chat")
    # the team chat is a real column next to the order details (was a 340 px strip)
    cw = op.eval_on_selector(".o-chat-card", "e => e.getBoundingClientRect().width"); gw = op.eval_on_selector(".o-grid", "e => e.getBoundingClientRect().width")
    ok(cw >= 380 and cw >= 0.4 * gw, "owner: the team chat column is wide (%d of %d px)" % (cw, gw))
    # a file with a long name (no spaces) in the chat must not push its bubble out of the thread
    long_name = os.path.join(OUT, "11-13_crown_v1_cement_gap_0.05_mm_final_exocad_export_v12_LONG.stl")
    open(long_name, "wb").write(b"solid long".ljust(84, b" "))
    op.set_input_files('#oChat [data-ochat="file"]', long_name); op.wait_for_timeout(300)
    op.fill("#oChat textarea", "file dài tên"); op.click("#oChat [type=submit]"); op.wait_for_timeout(1500)
    ok(op.query_selector("#oChat a.ochat-file") is not None and "_LONG" in text(op, "#oChat .ochat-list"), "owner sends a file with a long name in the chat")
    ok(op.eval_on_selector("#oChat .ochat-list", THREAD_FITS), "owner: nothing in the chat thread runs past the thread (long file name)")
    op.screenshot(path=f"{OUT}/order-page.png", full_page=True)

    # search
    old_style = "/".join([code_preview[0:2], code_preview[2:4], code_preview[4:6], code_preview[6:]])
    for q, why in [("chuang", "patient"), ("kiet", "doctor without accents"), (code_preview, "code"), (old_style, "code typed the old way"), (CODE, "customer code")]:
        op.goto(B + "/admin#orders"); op.wait_for_timeout(800)
        op.fill("#oQ", q); op.wait_for_timeout(1200)
        ok(code_preview in text(op, "#oList"), "list search by " + why)
    op.screenshot(path=f"{OUT}/orders-list.png", full_page=True)

    # the same case typed in again: the form says it is already in the book, with its code
    op.goto(B + f"/admin#orders/new?customer={CID}"); op.wait_for_timeout(1500)
    op.fill("#ofPatient", "chuang tsung wei"); op.wait_for_timeout(1300)
    ok(op.is_visible("#ofDupes") and code_preview in text(op, "#ofDupes"), "duplicate warning shows the existing code " + code_preview)
    op.fill('#ofItems [data-i="0"] [data-k="product"]', "CR"); op.wait_for_timeout(400)
    op.click("#ofSave"); op.wait_for_timeout(900)
    ok("#orders/new" in op.url and op.is_visible("#ofErr") and "đã có" in text(op, "#ofErr"), "saving is held back until confirmed")
    op.screenshot(path=f"{OUT}/order-dupe.png", full_page=True)
    op.check("#ofDupOk"); op.click("#ofSave"); op.wait_for_timeout(1800)
    m2 = re.search(r"#orders/(\d+)$", op.url)
    ok(m2 is not None, "after ticking 'this is another order' it is created")
    if m2: op.evaluate("async (id) => (await fetch('/api/orders/' + id, {method: 'DELETE', headers: {'X-D8': '1'}})).status", m2.group(1))

    # redo (R) and next stage (C): the code is shown before saving
    op.goto(B + f"/admin#orders/new?redo={OID}"); op.wait_for_timeout(1500)
    ok(text(op, "#ofKidCode") == code_preview + "R1", "redo form shows the new code " + text(op, "#ofKidCode"))
    op.click("#ofSave"); op.wait_for_timeout(1800)
    ok(code_preview + "R1" in op.inner_text(".pt-band"), "redo order " + code_preview + "R1")
    ok(code_preview in text(op, ".o-fam") and "R1" in text(op, ".o-fam"), "family of orders shown")
    op.goto(B + f"/admin#orders/new?continue={OID}"); op.wait_for_timeout(1500)
    ok(text(op, "#ofKidCode") == code_preview + "C2", "next-stage form shows " + text(op, "#ofKidCode") + " (stage 1 is the plain code)")
    op.click("#ofSave"); op.wait_for_timeout(1800)
    ok(code_preview + "C2" in op.inner_text(".pt-band") and "Giai đoạn 2" in op.inner_text(".pt-band"), "stage 2 order " + code_preview + "C2")

    # KTV: no money, workflow, chat
    kc = ctx(); kp = kc.new_page(); watch(kp, "ktv")
    kp.goto(B + "/admin"); kp.evaluate(SIGNIN, [KTV, "Minh KTV " + stamp]); kp.reload(); kp.wait_for_timeout(1500)
    nav = kp.eval_on_selector_all("#wsNav nav a > span:not(.count)", "els=>els.map(e=>e.textContent.trim())")
    ok("Báo cáo" not in nav and "Khách hàng" not in nav and "Mã hàng & bảng giá" not in nav, "KTV sidebar has no reports, customers or prices")
    ok(code_preview in text(kp, "#dash"), "KTV dashboard lists the order")
    try: kp.wait_for_selector(".sh-bell-n:not([hidden])", timeout=8000)
    except Exception: pass
    kp.click(".sh-bell-btn"); kp.wait_for_timeout(200)
    ok("đơn cần thiết kế" in text(kp, ".sh-bell-pop"), "KTV bell: orders to design")
    kp.click(".sh-bell-pop a"); kp.wait_for_timeout(1500)
    ok("#orders?designer=me" in kp.url and code_preview in text(kp, "#oList"), "bell item opens the KTV's orders to design")
    kp.goto(B + "/admin#orders/" + OID); kp.wait_for_timeout(1500)
    kb = kp.inner_text("#p-order")
    ok("US$" not in kb and "Thanh toán" not in kb and "Còn nợ" not in kb, "KTV order page shows no money")
    ok("CR" in kb and "IMC" in kb and "11-13" in kb, "KTV sees the items and teeth")
    kp.fill("#oChat textarea", "Anh ơi 12 cần dài thêm không?"); kp.click("#oChat [type=submit]"); kp.wait_for_timeout(900)
    # a design file in the chat
    # a finished design (binary STL, a crown-like mesh): QC and the owner rotate it on the order page
    import math, struct
    def crown_mesh(sub=3):
        t = (1 + 5 ** 0.5) / 2
        v = [(-1,t,0),(1,t,0),(-1,-t,0),(1,-t,0),(0,-1,t),(0,1,t),(0,-1,-t),(0,1,-t),(t,0,-1),(t,0,1),(-t,0,-1),(-t,0,1)]
        v = [tuple(c / math.sqrt(sum(x*x for x in p_)) for c in p_) for p_ in v]
        f = [(0,11,5),(0,5,1),(0,1,7),(0,7,10),(0,10,11),(1,5,9),(5,11,4),(11,10,2),(10,7,6),(7,1,8),(3,9,4),(3,4,2),(3,2,6),(3,6,8),(3,8,9),(4,9,5),(2,4,11),(6,2,10),(8,6,7),(9,8,1)]
        for _ in range(sub):
            nf, cache = [], {}
            def mid(a, b):
                k = tuple(sorted((a, b)))
                if k not in cache:
                    q = [(v[a][i] + v[b][i]) / 2 for i in range(3)]; l = math.sqrt(sum(x*x for x in q)); v.append(tuple(x / l for x in q)); cache[k] = len(v) - 1
                return cache[k]
            for a, b_, c in f:
                ab, bc, ca = mid(a, b_), mid(b_, c), mid(c, a); nf += [(a, ab, ca), (b_, bc, ab), (c, ca, bc), (ab, bc, ca)]
            f = nf
        return [(x * 5, y * 4.5, z * 3.6 + (0.6 if z > 0.3 else 0) * math.cos(x * 6) * math.cos(y * 6)) for x, y, z in v], f
    cv, cf = crown_mesh()
    stl = os.path.join(OUT, "11-13 crown v1.stl")
    with open(stl, "wb") as fh:
        fh.write(b"solid crown exocad".ljust(80, b" ")); fh.write(struct.pack("<I", len(cf)))
        for a, b_, c in cf: fh.write(struct.pack("<12fH", 0, 0, 0, *cv[a], *cv[b_], *cv[c], 0))
    STL_BYTES = open(stl, "rb").read()
    kp.set_input_files('#oChat [data-ochat="file"]', stl); kp.wait_for_timeout(300)
    ok("11-13 crown v1.stl" in text(kp, "#oChat .ochat-files"), "file waits under the message box")
    kp.fill("#oChat textarea", "File thiết kế v1"); kp.click("#oChat [type=submit]"); kp.wait_for_timeout(2000)
    ok(kp.query_selector("#oChat a.ochat-file") is not None and "11-13 crown v1.stl" in text(kp, "#oChat .ochat-list"), "file sent in the chat")
    ok(kp.eval_on_selector("#oChat .ochat-list", THREAD_FITS), "KTV: the chat thread fits its column")
    href = kp.eval_on_selector('#oChat a.ochat-file[download="11-13 crown v1.stl"]', "a => a.getAttribute('href')")
    fr = kp.request.get(B + href)
    ok(fr.status == 200 and fr.body() == STL_BYTES, "file downloads from the chat")
    kp.click('[data-act="submit"]'); kp.wait_for_timeout(1000)
    ok("Chờ QC" in kp.inner_text(".pt-band"), "KTV sends to QC")
    kp.screenshot(path=f"{OUT}/ktv-order.png", full_page=True)
    kp.goto(B + "/admin#reports"); kp.wait_for_timeout(600)
    ok(kp.is_visible("#p-denied"), "KTV cannot open reports")

    qcx = ctx(); qp = qcx.new_page(); watch(qp, "qc")
    qp.goto(B + "/admin"); qp.evaluate(SIGNIN, [QC, "Lan QC " + stamp]); qp.reload(); qp.wait_for_timeout(1500)
    qp.goto(B + "/admin#orders"); qp.wait_for_timeout(1200)
    ok(qp.query_selector("#oList .o-dot") is not None, "QC sees an unread chat mark")
    qp.goto(B + "/admin#orders/" + OID); qp.wait_for_timeout(1500)
    ok("12 cần dài thêm" in text(qp, "#oChat"), "QC reads the KTV message")
    try: qp.wait_for_selector("#o3d .v3d[data-tris]", timeout=15000)
    except Exception: pass
    ok(qp.get_attribute("#o3d .v3d", "data-tris") == str(len(cf)) and "tam giác" in text(qp, "#o3d .v3d-info"), "QC sees the design in 3D on the order page (" + str(qp.get_attribute("#o3d .v3d", "data-tris")) + " triangles)")
    PIX = """(sel) => { const c = document.querySelector(sel); if (!c) return 0; const t = document.createElement('canvas'); t.width = c.width; t.height = c.height; const x = t.getContext('2d'); x.drawImage(c, 0, 0);
      const d = x.getImageData(0, 0, t.width, t.height).data; let n = 0; for (let i = 3; i < d.length; i += 4) if (d[i] > 0) n++; return n; }"""
    px0 = qp.evaluate(PIX, "#o3d canvas")
    ok(px0 > 1000, "the model is drawn (" + str(px0) + " pixels)")
    qp.evaluate("document.querySelector('#o3d').scrollIntoView({block: 'center'})"); qp.wait_for_timeout(300)
    cb = qp.query_selector("#o3d canvas").bounding_box()
    shot0 = qp.query_selector("#o3d canvas").screenshot()
    qp.mouse.move(cb["x"] + cb["width"] / 2, cb["y"] + cb["height"] / 2); qp.mouse.down(); qp.mouse.move(cb["x"] + cb["width"] / 2 + 140, cb["y"] + cb["height"] / 2 + 60, steps=8); qp.mouse.up()
    qp.wait_for_timeout(200)
    ok(qp.query_selector("#o3d canvas").screenshot() != shot0, "drag rotates the model")
    px1 = qp.evaluate(PIX, "#o3d canvas")
    qp.mouse.wheel(0, -500); qp.wait_for_timeout(300)
    ok(qp.evaluate(PIX, "#o3d canvas") > px1, "wheel zooms in")
    qp.screenshot(path=f"{OUT}/order-3d.png", full_page=True)
    # several files in one view. A new file in the chat only adds its chip: the crown stays as it is (not loaded again)
    qp.evaluate("window.__c3 = document.querySelector('#o3d canvas')")
    ant = os.path.join(OUT, "11-13 antagonist.stl")
    with open(ant, "wb") as fh:
        fh.write(b"solid antagonist".ljust(80, b" ")); fh.write(struct.pack("<I", len(cf)))
        for a, b_, c in cf: fh.write(struct.pack("<12fH", 0, 0, 0, *[x - (9 if k == 2 else 0) for v_ in (cv[a], cv[b_], cv[c]) for k, x in enumerate(v_)], 0))
    qp.set_input_files('#oChat [data-ochat="file"]', ant); qp.wait_for_timeout(300)
    qp.fill("#oChat textarea", "Hàm đối diện"); qp.click("#oChat [type=submit]")
    try: qp.wait_for_selector("#o3d .v3d-chip[title='11-13 antagonist.stl']", timeout=15000)
    except Exception: pass
    ok(qp.query_selector("#o3d .v3d-chip[title='11-13 antagonist.stl']") is not None and qp.get_attribute("#o3d .v3d-chip.on", "title") == "11-13 crown v1.stl"
       and qp.evaluate("document.querySelector('#o3d canvas') === window.__c3"), "a new file in the chat adds its chip, the crown view is kept")
    ok(qp.is_visible("#o3d .v3d-mode") and qp.query_selector("#o3d .v3d-legend") is None, "3D: 'Từng file / Xem cùng lúc' switch, one file at first")
    qp.click("#o3d .v3d-mode label:has(input[value='multi'])"); qp.wait_for_timeout(200)
    qp.click("#o3d .v3d-chip[title='11-13 antagonist.stl']")
    try: qp.wait_for_selector("#o3d .v3d[data-models='2']", timeout=15000)
    except Exception: pass
    rows = qp.eval_on_selector_all("#o3d .v3d-legend li", "els => els.map((e) => [e.querySelector('b').textContent, e.querySelector('.v3d-dot').className])")
    ok(qp.get_attribute("#o3d .v3d", "data-models") == "2" and qp.get_attribute("#o3d .v3d", "data-tris") == str(2 * len(cf)) and len(rows) == 2
       and rows[0][1] != rows[1][1], "Xem cùng lúc: crown and antagonist in one view, each its own colour " + str(rows))
    ok(qp.evaluate("document.querySelector('#o3d canvas') === window.__c3"), "adding a file keeps the view (the crown is not loaded again)")
    s3 = qp.query_selector("#o3d canvas").screenshot()
    qp.click("#o3d .v3d-legend li:first-child [data-ghost]"); qp.wait_for_timeout(250)
    ok(qp.get_attribute("#o3d .v3d-legend li:first-child [data-ghost]", "aria-pressed") == "true" and qp.query_selector("#o3d canvas").screenshot() != s3, "'Mờ' makes a model see-through")
    qp.click(f"#o3d .v3d-chip[title='{os.path.basename(long_name)}']")
    try: qp.wait_for_selector("#o3d .v3d-legend li.err", timeout=15000)
    except Exception: pass
    ok(qp.query_selector("#o3d .v3d-legend li.err") is not None and qp.get_attribute("#o3d .v3d", "data-models") == "2", "a file that cannot be read shows its error, the others stay")
    qp.evaluate("document.querySelector('#o3d').scrollIntoView({block: 'center'})"); qp.wait_for_timeout(200)
    qp.screenshot(path=f"{OUT}/order-3d-multi.png", full_page=True)
    # the chat refreshes every 15 s: the view is left alone (it used to be rebuilt and every file loaded again)
    qp.wait_for_timeout(16000)
    ok(qp.evaluate("document.querySelector('#o3d canvas') === window.__c3") and qp.get_attribute("#o3d .v3d", "data-models") == "2", "a chat refresh leaves the 3D view as it is")
    qp.click("#o3d .v3d-mode label:has(input[value='one'])"); qp.wait_for_timeout(300)
    ok(qp.get_attribute("#o3d .v3d", "data-models") == "1" and qp.get_attribute("#o3d .v3d-chip.on", "title") == "11-13 antagonist.stl" and qp.query_selector("#o3d .v3d-legend") is None,
       "back to one file: the last one picked that shows")
    qp.click('#oChat .ochat-3d[data-name="11-13 crown v1.stl"]')
    try: qp.wait_for_selector(".v3d-overlay .v3d[data-tris]", timeout=15000)
    except Exception: pass
    ok(qp.is_visible(".v3d-overlay canvas"), "Xem 3D in the chat opens a large viewer")
    qp.screenshot(path=f"{OUT}/order-3d-full.png")
    qp.keyboard.press("Escape"); qp.wait_for_timeout(300)
    ok(qp.query_selector(".v3d-overlay") is None, "Esc closes it")
    qp.click('[data-act="fail"]'); qp.wait_for_timeout(500)
    ok("Chờ QC" in qp.inner_text(".pt-band"), "QC fail needs a note")
    qp.fill("#oNote", "Tiếp xúc 21 quá chặt"); qp.click('[data-act="fail"]'); qp.wait_for_timeout(1000)
    ok("QC trả sửa" in qp.inner_text(".pt-band"), "QC returns the order")
    try: qp.wait_for_selector("#o3d .v3d[data-tris]", timeout=10000)
    except Exception: pass
    ok(qp.evaluate("document.querySelector('#o3d canvas') === window.__c3"), "the order page redraws and keeps its 3D view (not loaded again)")
    kp.goto(B + "/admin#orders/" + OID); kp.wait_for_timeout(1500)
    ok("Tiếp xúc 21 quá chặt" in text(kp, "#oChat"), "KTV sees the QC note in the chat")
    kp.click('[data-act="submit"]'); kp.wait_for_timeout(900)
    qp.reload(); qp.wait_for_timeout(1500); qp.click('[data-act="pass"]'); qp.wait_for_timeout(1000)
    ok("Đã qua QC" in qp.inner_text(".pt-band"), "QC passes")

    # owner delivers and records the payment
    op.goto(B + "/admin#orders/" + OID); op.wait_for_timeout(1500)
    op.click('[data-act="deliver"]'); op.wait_for_timeout(1000)
    ok("Đã giao" in op.inner_text(".pt-band"), "owner delivers")
    op.click("#opAddBtn"); op.fill("#opRef", "PP-" + stamp); op.click("#opForm [type=submit]"); op.wait_for_timeout(1500)
    ok(re.search(r"Còn nợ\s*US\$0", op.inner_text("#p-order")) is not None, "payment recorded, nothing owed")

    # KTV pages: the team, one KTV's month, "Công của tôi", pay slips in Excel
    op.evaluate(API, ["/api/techrates", {"email": KTV, "currency": "VND", "from": "2026-01-01", "items": [{"product": "CR", "rate": 100000}, {"product": "IMC", "rate": 150000}]}])
    nav = op.eval_on_selector_all("#wsNav nav a > span:not(.count)", "els=>els.map(e=>e.textContent.trim())")
    ok("Kỹ thuật viên" in nav and "Công của tôi" not in nav, "owner sidebar: Kỹ thuật viên (not Công của tôi)")
    op.goto(B + "/admin#team"); op.wait_for_timeout(1500)
    row = op.query_selector(f'#tmBody tr[data-href="#team/{KTV.replace("@", "%40")}"]')
    ok(row is not None and "450.000" in row.inner_text(), "team page: KTV row with this month's pay 3×100,000 + 150,000")
    ok("1" in (row.query_selector('[data-l="QC trả về"]').inner_text() if row else ""), "team page: QC return counted")
    op.screenshot(path=f"{OUT}/team.png", full_page=True)
    with op.expect_download() as dl:
        op.click("#tmExport")
    allp = os.path.join(OUT, "phieu-cong-tat-ca.xlsx"); dl.value.save_as(allp)
    wb = openpyxl.load_workbook(allp)
    summ = list(wb["Tổng hợp"].iter_rows(values_only=True))
    ok(len(wb.sheetnames) >= 2 and any(r and r[1] == KTV and r[6] == 450000 for r in summ), "pay slips for all KTVs: summary + one sheet per KTV " + ", ".join(wb.sheetnames))
    if row: row.click()
    op.wait_for_timeout(1500)
    ok("#team/" in op.url and code_preview in text(op, "#tkBody") and "450.000" in text(op, "#tkBody .kpis"), "KTV page: this month's case and pay")
    ok(op.query_selector("#tkChart svg") is not None, "KTV page: 12-month chart")
    op.screenshot(path=f"{OUT}/tech.png", full_page=True)
    with op.expect_download() as dl:
        op.click("#tkExport")
    one = os.path.join(OUT, "phieu-cong.xlsx"); dl.value.save_as(one)
    wb = openpyxl.load_workbook(one)
    ok(wb.sheetnames == ["Phiếu công", "Chi tiết ca", "Theo mã hàng"], "pay slip sheets: " + ", ".join(wb.sheetnames))
    head = {r[0]: r[1] for r in wb["Phiếu công"].iter_rows(values_only=True) if r and r[0]}
    det = [r for r in wb["Chi tiết ca"].iter_rows(values_only=True) if r and r[0] == code_preview]
    ok(head.get("Email") == KTV and head.get("Tổng tiền công (VND)") == 450000 and head.get("Ca bị QC trả về") == 1, "pay slip head: KTV, total pay, QC returns " + json.dumps({k: str(v) for k, v in head.items()})[:220])
    ok(det and det[0][10] == 450000 and det[0][8] == 1, "pay slip line for " + code_preview)
    op.select_option("#tkBasis", "received"); op.wait_for_timeout(1200)
    ok("ca nhận tháng" in text(op, "#tkBody").lower(), "switch to received day")
    op.click('#p-tech [data-mon="-1"]'); op.wait_for_timeout(1200)
    ok(op.input_value("#tkMon") != op.evaluate("new Date(Date.now()+7*3600e3).toISOString().slice(0,7)"), "previous month")
    op.select_option("#tkBasis", "delivered"); op.select_option("#tkMon", op.evaluate("new Date(Date.now()+7*3600e3).toISOString().slice(0,7)")); op.wait_for_timeout(1200)

    kp.reload(); kp.wait_for_timeout(1500)
    nav = kp.eval_on_selector_all("#wsNav nav a > span:not(.count)", "els=>els.map(e=>e.textContent.trim())")
    ok("Công của tôi" in nav and "Kỹ thuật viên" not in nav, "KTV sidebar: Công của tôi, no team page")
    kp.goto(B + "/admin#me"); kp.wait_for_timeout(1500)
    kb = text(kp, "#tkBody")
    ok(code_preview in kb and "450.000" in kb and "US$" not in kb, "Công của tôi: own case and pay, no prices")
    kp.screenshot(path=f"{OUT}/ktv-me.png", full_page=True)
    with kp.expect_download() as dl:
        kp.click("#tkExport")
    ok(dl.value.suggested_filename.startswith("D8-phieu-cong-"), "KTV downloads their own pay slip")
    kp.goto(B + "/admin#team"); kp.wait_for_timeout(1000)
    ok(kp.is_visible("#p-denied"), "KTV cannot open the team page")
    kp.goto(B + "/admin#team/" + QC.replace("@", "%40")); kp.wait_for_timeout(1000)
    ok(kp.is_visible("#p-denied"), "KTV cannot open someone else's page")
    qp.reload(); qp.wait_for_timeout(1500)
    nav = qp.eval_on_selector_all("#wsNav nav a > span:not(.count)", "els=>els.map(e=>e.textContent.trim())")
    ok("Công của tôi" not in nav and "Kỹ thuật viên" not in nav, "QC sidebar: no KTV pages")

    # reports + Excel export
    op.goto(B + "/admin#reports"); op.wait_for_timeout(1800)
    ok(op.query_selector("#rpMoney svg path, #rpMoney svg rect:not(.hit)") is not None and op.query_selector("#rpCases svg") is not None, "report charts drawn")
    hit = op.query_selector_all("#rpCases .hit")
    if hit:
        hit[-1].hover(); op.wait_for_timeout(300)
        ok(op.is_visible("#rpCases .chart-tip"), "chart tooltip on hover")
    op.screenshot(path=f"{OUT}/reports.png", full_page=True)
    with op.expect_download() as dl:
        op.click("#rpExport")
    path = os.path.join(OUT, "export.xlsx"); dl.value.save_as(path)
    wb = openpyxl.load_workbook(path)
    ok(wb.sheetnames[:3] == ["Tổng hợp tháng", "Đơn hàng", "Hạng mục"] and "Theo khách" in wb.sheetnames and "Theo KTV" in wb.sheetnames, "export sheets: " + ", ".join(wb.sheetnames))
    rows = list(wb["Đơn hàng"].iter_rows(values_only=True))
    hit_row = [r for r in rows if r and r[0] == code_preview]
    ok(hit_row and hit_row[0][6] == "Chuang Tsung Wei" and 225 in hit_row[0], "exported order row has patient and total")
    ok(isinstance(rows[1][1], __import__("datetime").datetime), "dates exported as real Excel dates")

    # Excel import of old orders, from the template
    op.goto(B + "/admin#import/orders"); op.wait_for_timeout(800)
    with op.expect_download() as dl:
        op.click("#imTpl")
    tpl = os.path.join(OUT, "mau-don.xlsx"); dl.value.save_as(tpl)
    wb = openpyxl.load_workbook(tpl); ws = wb["Đơn hàng"]
    ws.delete_rows(2, ws.max_row)
    for r in [["B1", "05/08/2026", CODE, "Lux", "Dr Kiệt", "Old patient " + stamp, "", "CR", "11 21", "", "", "", "", "Đã giao", "", "", "", "x"],
              ["B1", "05/08/2026", CODE, "", "", "", "", "IMC", "", 1, 40],
              ["", "06/08/2026", CODE, "", "", "Second old " + stamp, "", "CR", "", 1, "", "", "Miễn phí"],
              ["", "07/08/2026", "NOPE" + stamp, "", "", "Bad", "", "CR"]]:
        ws.append(r)
    filled = os.path.join(OUT, "don-cu.xlsx"); wb.save(filled)
    op.set_input_files("#imFile", filled); op.wait_for_timeout(2500)
    prev = text(op, "#imBody")
    ok("không có khách mã" in prev.lower() or "NOPE" in prev, "preview flags the unknown customer")
    ok(op.is_enabled("#imGo") and "2" in op.inner_text("#imGo"), "2 valid orders ready: " + op.inner_text("#imGo"))
    op.screenshot(path=f"{OUT}/import-preview.png", full_page=True)
    op.click("#imGo"); op.wait_for_timeout(2500)
    ok("Đã nhập 2" in text(op, "#imDone"), "imported 2 old orders")
    st, lst = op.evaluate(API, [f"/api/orders?customer={CID}&status=&limit=50&q=old"])
    olds = lst.get("items", [])
    o1 = next((x for x in olds if x["patient"] == "Old patient " + stamp), None)
    ok(o1 and o1["day"] == "2026-08-05" and o1["total"] == 110 and o1["paid"] == 110 and o1["units"] == 3, "imported order: 2×CR(35) + IMC 40, paid: " + json.dumps(o1)[:160] if o1 else "imported order missing")

    # a KTV's own sheet ("GC - Case Detail": no customer column, doctors, product names, US dates, Progress)
    CODE2 = "H" + stamp
    op.evaluate(API, ["/api/customers", {"code": CODE2, "name": "Hana " + stamp, "kind": "clinic", "currency": "USD"}])
    gc = openpyxl.Workbook(); gs = gc.active; gs.title = "GC - Case Detail"
    gs.append(["Case_ID", "Created_Date", "Doctor_Name", "Patient_Name", "Product_Name", "Unit", "Instruction_Note", "Progress"])
    for r in [["GC" + stamp + "1", "10/14/2026", "Dr Gene", "GC Alpha " + stamp, "Crown (full anatomy)", 2, "Light contact", "Completed"],
              ["GC" + stamp + "1", "10/14/2026", "Dr Gene", "GC Alpha " + stamp, "Implant crown", 1, "", "Completed"],
              ["GC" + stamp + "2", "10/15/2026", "Dr Gene", "GC Beta " + stamp, "Night guard XYZ", 1, "", "Pending"],
              ["GC" + stamp + "3", __import__("datetime").datetime(2026, 10, 16), "Dr Gene", "GC Gamma " + stamp, "Zirconia crown special", 1, "", "Completed"],
              ["GC" + stamp + "4", "10/16/2026", "Dr. Hana " + stamp, "GC Delta " + stamp, "crown (full anatomy)", 1, "", "Completed"]]:
        gs.append(r)
    gcf = os.path.join(OUT, "GC - Case Detail.xlsx"); gc.save(gcf)
    op.goto(B + "/admin#import/orders"); op.wait_for_timeout(1000)
    op.set_input_files("#imFile", gcf); op.wait_for_timeout(2500)
    docs = op.eval_on_selector_all("#imCmapBody .im-map-row span", "els => els.map((e) => e.textContent)")
    ok(docs == ["Dr Gene"], "doctor named like a customer is matched by itself; the other is asked: " + ", ".join(docs))
    names = op.eval_on_selector_all("#imMapBody .im-map-row span", "els => els.map((e) => e.textContent)")
    ok(sorted(names) == ["Night guard XYZ", "Zirconia crown special"], "only names not in the catalogue need matching: " + ", ".join(names))
    ok("nhập 1" in op.inner_text("#imGo").lower(), "the case of the matched doctor is ready: " + op.inner_text("#imGo"))
    op.select_option("#imKtv", KTV); op.wait_for_timeout(1500)
    op.select_option('#imCmapBody [data-map="gene"]', CODE); op.wait_for_timeout(2000)
    ok("nhập 2" in op.inner_text("#imGo").lower(), "after matching Dr Gene: 2 ready: " + op.inner_text("#imGo"))
    op.select_option('#imMapBody [data-map="night guard xyz"]', "NG"); op.wait_for_timeout(1800)
    op.select_option('#imMapBody [data-map="zirconia crown special"]', "CR"); op.wait_for_timeout(1800)
    ok("nhập 4" in op.inner_text("#imGo").lower() and op.is_visible("#imMap") and op.is_visible("#imCmap"), "after matching names: 4 orders ready: " + op.inner_text("#imGo"))
    ok(op.is_visible("#imZero") and "thành tiền 0" in text(op, "#imZero"), "warns about orders with no price")
    op.screenshot(path=f"{OUT}/import-gc.png", full_page=True)
    op.click("#imGo"); op.wait_for_timeout(2500)
    ok("Đã nhập 4" in text(op, "#imDone"), "imported 4 orders from the KTV's sheet")
    st, lst = op.evaluate(API, ["/api/orders?status=&limit=100&q=gc"])
    gcs = {x["patient"].split(" ")[1]: x for x in lst.get("items", []) if x["patient"].startswith("GC ") and x["patient"].endswith(stamp)}
    a, bt, g, d = gcs.get("Alpha"), gcs.get("Beta"), gcs.get("Gamma"), gcs.get("Delta")
    ok(a and a["day"] == "2026-10-14" and a["units"] == 3 and a["status"] == "delivered" and a["ref"] == "GC" + stamp + "1" and a["customer_code"] == CODE and a["designer"] == KTV,
       "lines with the same Case_ID are one order, US date, customer and KTV: " + json.dumps(a)[:240] if a else "GC order missing")
    st, ad = op.evaluate(API, ["/api/orders/" + str(a["id"])]) if a else (0, {})
    its = sorted((x["product"], x["qty"]) for x in ad.get("items", []))
    ok(its == [("CR", 2), ("IMC", 1)] and "Light contact" in (ad.get("order", {}).get("notes") or ""), "items by product name: " + str(its))
    ok(bt and bt["status"] == "design" and bt["day"] == "2026-10-15", "Pending → in design")
    ok(g and g["day"] == "2026-10-16", "real Excel date cell read")
    ok(d and d["customer_code"] == CODE2, "Dr. Hana's case went to customer Hana")

    # 3D link with a preview picture, attached to the order
    op.goto(B + "/admin#links"); op.wait_for_timeout(900)
    op.fill("#ogOrder", code_preview); op.dispatch_event("#ogOrder", "change"); op.wait_for_timeout(1200)
    ok("CHUANG TSUNG WEI" in op.input_value("#label") and code_preview in op.input_value("#ogDesc"), "order code fills the preview title and description")
    op.set_input_files("#file", EXO); op.wait_for_timeout(400)
    op.click("#create"); op.wait_for_timeout(2500)
    for _ in range(30):
        if "Đã lưu ảnh" in text(op, "#snapSay") or "mặc định" in text(op, "#snapSay"): break
        op.wait_for_timeout(500)
    say = text(op, "#snapSay")
    ok("Đã lưu ảnh" in say, "preview picture captured from the 3D view: " + say)
    link = op.inner_text("#resUrl"); lid = link.rsplit("/", 1)[-1]
    op.screenshot(path=f"{OUT}/link-preview.png", full_page=True)
    r = op.request.get(B + "/og/" + lid + ".jpg"); default = op.request.get(B + "/brand/og-3d.jpg")
    ok(r.status == 200 and r.headers["content-type"] == "image/jpeg" and r.body() != default.body(), "link preview picture is the captured design")
    ok(len(r.body()) < 300 * 1024, "captured preview picture under 300 KB, so WhatsApp shows it (" + str(len(r.body())) + " bytes)")
    op.click("#snapDef")
    for _ in range(20):
        if "mặc định" in text(op, "#snapSay"): break
        op.wait_for_timeout(300)
    r2 = op.request.get(B + "/og/" + lid + ".jpg")
    ok(r2.body() == default.body() and len(default.body()) < 300 * 1024, "'Dùng ảnh mặc định' gives the link the D.EIGHT default picture")
    page = op.request.get(B + "/v/" + lid).text()
    ok('og:title" content="Patient: CHUANG TSUNG WEI' in page, "doctor link carries the preview title")
    st, od = op.evaluate(API, ["/api/orders/" + OID])
    ok(lid in od["order"]["previews"], "3D link attached to the order")

    # phone layout
    mc = ctx({"width": 390, "height": 844}, True); mp = mc.new_page(); watch(mp, "mobile")
    mp.goto(B + "/admin"); mp.evaluate(SIGNIN, ["d.eightdesigncenter@gmail.com", "Danny Do"]); mp.reload(); mp.wait_for_timeout(1200)
    # quick buttons: in the band on a wide screen, under it on a phone (no layout shift after load)
    ok(mp.is_visible("#dash .dash-quick a[href='#orders/new']") and not mp.is_visible("#bandRight .home-q"), "phone: quick order buttons under the band")
    op.evaluate("location.hash = '#home'")
    op.wait_for_selector("#dash .card", timeout=10000); op.wait_for_timeout(600)
    ok(op.is_visible("#bandRight .home-q a[href='#orders/new']") and not op.is_visible("#dash .dash-quick"), "desktop: quick order buttons in the band")
    mp.goto(B + "/admin#orders/" + OID); mp.wait_for_timeout(1500)
    rows = mp.eval_on_selector_all(".o-items tbody tr", "els => els.map((e) => e.getBoundingClientRect().width)")
    ok(rows and all(w <= 360 for w in rows) and mp.eval_on_selector(".o-items tbody tr td", "e => e.getBoundingClientRect().width") > 250, "phone: order items as full-width rows")
    ok(mp.eval_on_selector("#oChat .ochat-list", THREAD_FITS) and mp.eval_on_selector(".o-chat-card", "e => e.getBoundingClientRect().width") <= 360, "phone: the chat thread fits the screen")
    for h in ["orders", "orders/" + OID, "orders/new", "reports", "catalog", "cust/" + CID, "import/orders", "team", "team/" + KTV.replace("@", "%40")]:
        mp.goto(B + "/admin#" + h); mp.wait_for_timeout(1500)
        w = mp.evaluate("document.documentElement.scrollWidth")
        ok(w <= 390, "phone: no sideways scroll on #" + h + " (" + str(w) + ")")
        mp.screenshot(path=f"{OUT}/m-{h.replace('/', '-')}.png", full_page=True)
    b.close()
for e in errors: print("ERR", e)
ok(not errors, "no page errors")
print(("%d FAILED" % len(FAILS)) if FAILS else "ALL PASSED")
sys.exit(1 if FAILS else 0)
